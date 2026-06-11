from __future__ import annotations

from dataclasses import asdict, dataclass
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill


DEFAULT_BATCH_COLUMNS = [
    "order_name",
    "quantity",
    "material_price_per_kg",
    "model_weight_g",
    "material_waste_rate",
    "print_hours",
    "printer_power_w",
    "electricity_price_per_kwh",
    "machine_purchase_price",
    "machine_lifetime_hours",
    "maintenance_cost_per_hour",
    "labor_hours",
    "labor_rate_per_hour",
    "packaging_cost",
    "other_fixed_cost",
    "failure_rate",
    "platform_fee_rate",
    "target_profit_margin",
    "override_unit_price",
]


@dataclass(frozen=True)
class PrintJobInput:
    order_name: str = "3D print order"
    quantity: int = 1
    material_price_per_kg: float = 80.0
    model_weight_g: float = 80.0
    material_waste_rate: float = 0.08
    print_hours: float = 4.0
    printer_power_w: float = 120.0
    electricity_price_per_kwh: float = 0.8
    machine_purchase_price: float = 2500.0
    machine_lifetime_hours: float = 2500.0
    maintenance_cost_per_hour: float = 0.5
    labor_hours: float = 0.5
    labor_rate_per_hour: float = 35.0
    packaging_cost: float = 3.0
    other_fixed_cost: float = 0.0
    failure_rate: float = 0.08
    platform_fee_rate: float = 0.06
    target_profit_margin: float = 0.30
    override_unit_price: float | None = None


@dataclass(frozen=True)
class QuoteResult:
    order_name: str
    quantity: int
    material_cost: float
    electricity_cost: float
    machine_cost: float
    labor_cost: float
    packaging_cost: float
    other_fixed_cost: float
    base_cost_before_failure: float
    failure_allowance: float
    cost_after_failure: float
    platform_fee: float
    unit_total_cost: float
    suggested_unit_price: float
    used_unit_price: float
    unit_profit: float
    actual_profit_margin: float
    profit_per_print_hour: float
    order_total_cost: float
    order_revenue: float
    order_profit: float

    def to_display_dict(self) -> dict[str, Any]:
        return {
            "订单名称": self.order_name,
            "数量": self.quantity,
            "材料成本": self.material_cost,
            "电费成本": self.electricity_cost,
            "机器折旧成本": self.machine_cost,
            "人工成本": self.labor_cost,
            "包装成本": self.packaging_cost,
            "其他固定成本": self.other_fixed_cost,
            "失败率摊销成本": self.failure_allowance,
            "平台手续费": self.platform_fee,
            "总成本 / 件": self.unit_total_cost,
            "建议售价 / 件": self.suggested_unit_price,
            "采用售价 / 件": self.used_unit_price,
            "利润 / 件": self.unit_profit,
            "利润率": self.actual_profit_margin,
            "每小时收益": self.profit_per_print_hour,
            "订单总成本": self.order_total_cost,
            "订单销售额": self.order_revenue,
            "每单净利润": self.order_profit,
        }


def _rate(value: float, field_name: str) -> float:
    value = float(value)
    if value < 0:
        raise ValueError(f"{field_name} 不能为负数。")
    if value >= 1:
        raise ValueError(f"{field_name} 必须小于 100%。")
    return value


def _non_negative(value: float, field_name: str) -> float:
    value = float(value)
    if value < 0:
        raise ValueError(f"{field_name} 不能为负数。")
    return value


def _positive(value: float, field_name: str) -> float:
    value = float(value)
    if value <= 0:
        raise ValueError(f"{field_name} 必须大于 0。")
    return value


def _clean_quantity(quantity: int | float) -> int:
    quantity = int(quantity)
    if quantity <= 0:
        raise ValueError("订单数量必须大于 0。")
    return quantity


def calculate_single_quote(job: PrintJobInput) -> QuoteResult:
    quantity = _clean_quantity(job.quantity)
    material_price_per_kg = _non_negative(job.material_price_per_kg, "材料价格")
    model_weight_g = _non_negative(job.model_weight_g, "模型重量")
    material_waste_rate = _rate(job.material_waste_rate, "材料损耗率")
    print_hours = _positive(job.print_hours, "打印时间")
    printer_power_w = _non_negative(job.printer_power_w, "打印机功率")
    electricity_price_per_kwh = _non_negative(job.electricity_price_per_kwh, "电价")
    machine_purchase_price = _non_negative(job.machine_purchase_price, "设备购置成本")
    machine_lifetime_hours = _positive(job.machine_lifetime_hours, "设备折旧寿命")
    maintenance_cost_per_hour = _non_negative(job.maintenance_cost_per_hour, "维护损耗")
    labor_hours = _non_negative(job.labor_hours, "人工时间")
    labor_rate_per_hour = _non_negative(job.labor_rate_per_hour, "人工单价")
    packaging_cost = _non_negative(job.packaging_cost, "包装成本")
    other_fixed_cost = _non_negative(job.other_fixed_cost, "其他固定成本")
    failure_rate = _rate(job.failure_rate, "失败率")
    platform_fee_rate = _rate(job.platform_fee_rate, "平台佣金率")
    target_profit_margin = _rate(job.target_profit_margin, "目标利润率")

    material_cost = (material_price_per_kg / 1000) * model_weight_g * (1 + material_waste_rate)
    electricity_cost = (printer_power_w / 1000) * print_hours * electricity_price_per_kwh
    machine_hourly_cost = (machine_purchase_price / machine_lifetime_hours) + maintenance_cost_per_hour
    machine_cost = machine_hourly_cost * print_hours
    labor_cost = labor_hours * labor_rate_per_hour

    base_cost = (
        material_cost
        + electricity_cost
        + machine_cost
        + labor_cost
        + packaging_cost
        + other_fixed_cost
    )
    cost_after_failure = base_cost / (1 - failure_rate)
    failure_allowance = cost_after_failure - base_cost

    price_denominator = 1 - platform_fee_rate - target_profit_margin
    if price_denominator <= 0:
        raise ValueError("平台佣金率 + 目标利润率必须小于 100%，否则无法反推出建议售价。")

    suggested_unit_price = cost_after_failure / price_denominator
    used_unit_price = (
        _positive(job.override_unit_price, "手动售价")
        if job.override_unit_price not in (None, 0)
        else suggested_unit_price
    )
    platform_fee = used_unit_price * platform_fee_rate
    unit_total_cost = cost_after_failure + platform_fee
    unit_profit = used_unit_price - unit_total_cost
    actual_profit_margin = unit_profit / used_unit_price if used_unit_price else 0.0
    profit_per_print_hour = unit_profit / print_hours if print_hours else 0.0

    return QuoteResult(
        order_name=job.order_name or "3D print order",
        quantity=quantity,
        material_cost=round(material_cost, 4),
        electricity_cost=round(electricity_cost, 4),
        machine_cost=round(machine_cost, 4),
        labor_cost=round(labor_cost, 4),
        packaging_cost=round(packaging_cost, 4),
        other_fixed_cost=round(other_fixed_cost, 4),
        base_cost_before_failure=round(base_cost, 4),
        failure_allowance=round(failure_allowance, 4),
        cost_after_failure=round(cost_after_failure, 4),
        platform_fee=round(platform_fee, 4),
        unit_total_cost=round(unit_total_cost, 4),
        suggested_unit_price=round(suggested_unit_price, 4),
        used_unit_price=round(used_unit_price, 4),
        unit_profit=round(unit_profit, 4),
        actual_profit_margin=round(actual_profit_margin, 4),
        profit_per_print_hour=round(profit_per_print_hour, 4),
        order_total_cost=round(unit_total_cost * quantity, 4),
        order_revenue=round(used_unit_price * quantity, 4),
        order_profit=round(unit_profit * quantity, 4),
    )


def normalize_order_dataframe(df: pd.DataFrame, defaults: PrintJobInput | None = None) -> pd.DataFrame:
    if df.empty:
        raise ValueError("订单表为空。")

    defaults = defaults or PrintJobInput()
    default_values = asdict(defaults)
    normalized = df.copy()
    normalized.columns = [str(column).strip() for column in normalized.columns]

    for column in DEFAULT_BATCH_COLUMNS:
        if column not in normalized.columns:
            normalized[column] = default_values.get(column)

    normalized = normalized[DEFAULT_BATCH_COLUMNS]
    normalized["order_name"] = normalized["order_name"].fillna("3D print order").astype(str)

    numeric_columns = [column for column in DEFAULT_BATCH_COLUMNS if column != "order_name"]
    for column in numeric_columns:
        normalized[column] = pd.to_numeric(normalized[column], errors="coerce")
        fallback = default_values.get(column, 0)
        normalized[column] = normalized[column].fillna(0 if fallback is None else fallback)

    normalized["override_unit_price"] = normalized["override_unit_price"].replace({0: None})
    return normalized


def calculate_batch_quotes(df: pd.DataFrame) -> pd.DataFrame:
    results: list[dict[str, Any]] = []

    for _, row in df.iterrows():
        job = PrintJobInput(
            order_name=str(row["order_name"]),
            quantity=int(row["quantity"]),
            material_price_per_kg=float(row["material_price_per_kg"]),
            model_weight_g=float(row["model_weight_g"]),
            material_waste_rate=float(row["material_waste_rate"]),
            print_hours=float(row["print_hours"]),
            printer_power_w=float(row["printer_power_w"]),
            electricity_price_per_kwh=float(row["electricity_price_per_kwh"]),
            machine_purchase_price=float(row["machine_purchase_price"]),
            machine_lifetime_hours=float(row["machine_lifetime_hours"]),
            maintenance_cost_per_hour=float(row["maintenance_cost_per_hour"]),
            labor_hours=float(row["labor_hours"]),
            labor_rate_per_hour=float(row["labor_rate_per_hour"]),
            packaging_cost=float(row["packaging_cost"]),
            other_fixed_cost=float(row["other_fixed_cost"]),
            failure_rate=float(row["failure_rate"]),
            platform_fee_rate=float(row["platform_fee_rate"]),
            target_profit_margin=float(row["target_profit_margin"]),
            override_unit_price=(
                None
                if pd.isna(row["override_unit_price"]) or float(row["override_unit_price"]) <= 0
                else float(row["override_unit_price"])
            ),
        )
        results.append(asdict(calculate_single_quote(job)))

    return pd.DataFrame(results)


def dataframe_to_excel_bytes(df: pd.DataFrame, sheet_name: str = "Results") -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        header_fill = PatternFill(fill_type="solid", fgColor="E8F1FF")
        header_font = Font(bold=True)

        for col_idx, column in enumerate(df.columns):
            cell = worksheet.cell(row=1, column=col_idx + 1)
            cell.fill = header_fill
            cell.font = header_font

            header = str(column)
            width = min(max(len(header) + 4, 12), 28)
            worksheet.column_dimensions[cell.column_letter].width = width

            for row_idx in range(2, len(df) + 2):
                data_cell = worksheet.cell(row=row_idx, column=col_idx + 1)
                if isinstance(data_cell.value, (int, float)):
                    data_cell.number_format = "0.00"

        worksheet.freeze_panes = "A2"
    return output.getvalue()


def dataframe_to_csv_bytes(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False).encode("utf-8-sig")
